import io
from typing import List, Dict
from datetime import datetime


def _fmt(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h > 0 else f"{m:02d}:{s:02d}"


def export_txt(
    segments: List[Dict],
    processed: str = "",
    action_label: str = "",
) -> str:
    lines: List[str] = []

    if segments:
        lines += ["TRANSCRIPCIÓN", "=" * 60,
                  f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}", ""]

        current_speaker = None
        for seg in segments:
            speaker = seg.get("speaker", "Hablante 1")
            text = seg.get("text", "").strip()
            start = seg.get("start", 0)
            if not text:
                continue
            if speaker != current_speaker:
                if current_speaker is not None:
                    lines.append("")
                lines.append(f"── {speaker} ──")
                current_speaker = speaker
            lines.append(f"[{_fmt(start)}] {text}")

    if processed and action_label:
        lines += ["", "", action_label.upper(), "=" * 60, "", processed]

    return "\n".join(lines)


def export_docx(
    segments: List[Dict],
    processed: str = "",
    action_label: str = "",
    filename: str = "transcripcion",
    show_timestamps: bool = True,
) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2.5)

    t = doc.add_heading("Transcripción", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(filename)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(12)
    sub.runs[0].font.color.rgb = RGBColor(120, 120, 120)

    d = doc.add_paragraph(datetime.now().strftime("%d/%m/%Y %H:%M"))
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.runs[0].font.size = Pt(10)
    d.runs[0].font.color.rgb = RGBColor(150, 150, 150)
    doc.add_paragraph()

    if segments:
        doc.add_heading("Transcripción", 1)
        current_speaker = None
        for seg in segments:
            speaker = seg.get("speaker", "Hablante 1")
            text = seg.get("text", "").strip()
            start = seg.get("start", 0)
            if not text:
                continue

            if speaker != current_speaker:
                sp = doc.add_paragraph()
                run = sp.add_run(speaker)
                run.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(201, 123, 95)
                current_speaker = speaker

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            if show_timestamps:
                ts = p.add_run(f"[{_fmt(start)}] ")
                ts.font.size = Pt(9)
                ts.font.color.rgb = RGBColor(160, 160, 160)
            txt = p.add_run(text)
            txt.font.size = Pt(11)

    if processed and action_label:
        doc.add_page_break()
        doc.add_heading(action_label, 1)
        for line in processed.split("\n"):
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
